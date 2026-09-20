#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
High-Fidelity DOCX to Markdown Converter with Universal Formula Mapping & Layout Recovery.

Features:
- Pure Standard Markdown: Strictly zero HTML tags (<p>, <b>, </p>, <br>).
- Dynamic Numbering Tracker: Parses word/numbering.xml to recover w:numPr heading and list prefixes.
- Dual Formula Mapping: Supports both equation tag mapping ("5.1-1") and OLE index mapping ("19" or 19).
- Safe Lambda Regex: Avoids Python re.sub "bad escape \\c" crashes on LaTeX commands.
- Side-by-Side Multi-Figure Unbundling: Detects independent captions and splits multi-image paragraphs cleanly.
- Ghost Caption Stripping: Cleans trailing captions leaked into headings or body text via zero-width assertions.
- DrawingML Deduplication: Prevents double-counting images nested under w:r.
- Relative Image Paths: Automatically calculates relative paths from output_md.
- Legacy .doc Support: Auto-detects and converts Word 97-2003 binary .doc via headless LibreOffice.
- CJK-Western Spacing: Automatic pangu typography micro-spacing between Chinese and Latin/numbers.
- Continuous Bold Merging: Merges fragmented consecutive bold runs into cohesive emphasis blocks.
- Outline & List Guard: Protects numbered question lists from being arbitrarily promoted to H4 headings.
"""

import os
import sys
import io
import re
import json
import argparse
import zipfile
import subprocess
import xml.etree.ElementTree as ET
import docx

# Configure stdout and stderr for UTF-8 encoding across Windows and Linux
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'o': 'urn:schemas-microsoft-com:office:office',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'v': 'urn:schemas-microsoft-com:vml',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'
}

PUA_GLYPH_MAP = {
    '\uf0b0': '^{\\circ}',
    '\uf0ce': ' \\in ',
    '\uf03d': ' = ',
    '\uf0e5': ' \\sum ',
    '\uf0a3': ' \\le ',
    '\uf0b3': ' \\ge ',
    '\uf0e9': ' \\lceil ',
    '\uf0f9': ' \\rceil ',
    '\uf0ea': ' \\lfloor ',
    '\uf0fa': ' \\rfloor ',
    '\uf02b': ' + ',
    '\uf02d': ' - ',
    '\uf0b4': ' \\times ',
    '\uf022': ' \\forall ',
    '\uf071': ' \\theta ',
    '\uf072': ' \\rho ',
    '\uf06c': ' \\lambda ',
    '\uf06a': ' \\varphi ',
    '\uf065': ' \\varepsilon ',
    '\uf070': ' \\pi ',
    '\uf0b6': ' \\partial ',
    '\uf0b9': ' \\neq ',
    '\uf044': ' \\Delta ',
    '\uf0bc': ' \\cdots ',
    '\uf051': ' \\boldsymbol{\\Theta} ',
    '\uf028': '(',
    '\uf029': ')',
    '\uf0ef': '',
    '\uf0ec': '',
    '\uf0ed': '',
    '\uf0ee': '',
    '\uf0eb': '',
    '\uf0fb': '',
    '\uff1d': ' = ',
    '\u00d7': ' \\times ',
    '\u00b0': '^{\\circ}',
}

def clean_inline_text(text):
    """Normalize whitespace, strip HTML tags, fix math spacing, and remove ghost captions."""
    if not text:
        return ""
    text = text.replace('\xa0', ' ')
    # Map PUA characters
    for glyph, repl in PUA_GLYPH_MAP.items():
        if glyph in text:
            text = text.replace(glyph, repl)
    text = "".join(c for c in text if not (0xE000 <= ord(c) <= 0xF8FF))

    # Strip HTML tags
    text = re.sub(r'<p\s+align=[^>]+><b>(.*?)</b></p>', r'**\1**', text, flags=re.DOTALL)
    text = re.sub(r'</?[a-zA-Z0-9]+(?:\s+[^>]*)?>', '', text)

    # Only remove trailing inline figure captions if this is body/heading text, not a standalone caption
    t_strip = text.strip()
    if not (t_strip.startswith("图") or t_strip.startswith("**图") or t_strip.startswith("Figure")):
        text = re.sub(r'[ \t]{2,}图\s*[\d\.\-]+[^\n]+$', '', text)
        text = re.sub(r'(?<=\*\*)\s*图\s*[\d\.\-]+[^\n]+$', '', text)

    # Normalize spaces
    text = re.sub(r'[ \t]+', ' ', text)
    # Fix repeated math dollar spaces: $ S_i $ -> $S_i$
    text = re.sub(r'\$\s+([^\$]+?)\s+\$', r'$\1$', text)
    text = text.replace('$ $', ' ').replace('$$', '$')
    # Merge adjacent bold spans: **A** **B** -> **A B**, **A****B** -> **AB**
    while re.search(r'\*\*(.*?)\*\*\s*\*\*(.*?)\*\*', text):
        text = re.sub(r'\*\*(.*?)\*\*\s*\*\*(.*?)\*\*', lambda m: f"**{m.group(1).strip()} {m.group(2).strip()}**", text)
    while re.search(r'\*\*(.*?)\*\*\*\*(.*?)\*\*', text):
        text = re.sub(r'\*\*(.*?)\*\*\*\*(.*?)\*\*', lambda m: f"**{m.group(1)}{m.group(2)}**", text)
    text = text.replace('****', '')
    # Normalize CJK and Western/Number spacing, protecting math/links/images
    parts = re.split(r"(\$[^\$]+?\$|!\[.*?\]\(.*?\)|\*\*.*?\*\*)", text)
    for idx in range(len(parts)):
        p = parts[idx]
        if not p.startswith(("$", "![", "**")):
            p = re.sub(r"([\u4e00-\u9fa5])([a-zA-Z0-9])", r"\1 \2", p)
            p = re.sub(r"([a-zA-Z0-9])([\u4e00-\u9fa5])", r"\1 \2", p)
            p = re.sub(r"（(\d+)）\s+", r"（\1）", p)
            parts[idx] = p
        elif p.startswith("**") and p.endswith("**"):
            inner = p[2:-2]
            inner = re.sub(r"([\u4e00-\u9fa5])([a-zA-Z0-9])", r"\1 \2", inner)
            inner = re.sub(r"([a-zA-Z0-9])([\u4e00-\u9fa5])", r"\1 \2", inner)
            parts[idx] = f"**{inner}**"
    text = "".join(parts)
    text = re.sub(r"([\u4e00-\u9fa5])(\$[^\$]+?\$)", lambda m: f"{m.group(1)} {m.group(2)}", text)
    text = re.sub(r"(\$[^\$]+?\$)([\u4e00-\u9fa5])", lambda m: f"{m.group(1)} {m.group(2)}", text)
    # Clean redundant spaces before punctuation
    text = re.sub(r'\s+([，。！？；：、\),])', r'\1', text)
    text = re.sub(r'([\(\(])\s+', r'\1', text)
    return text.strip()

def build_numbering_tracker(docx_path):
    """Parse word/numbering.xml to recover dynamic list & heading numbering (w:numPr)."""
    try:
        with zipfile.ZipFile(docx_path) as z:
            if 'word/numbering.xml' not in z.namelist():
                return lambda numPr: None
            num_xml = z.read('word/numbering.xml')
        root = ET.fromstring(num_xml)
    except Exception:
        return lambda numPr: None

    abstract_nums = {}
    for ab in root.findall('w:abstractNum', NS):
        ab_id = ab.attrib.get(f'{{{NS["w"]}}}abstractNumId')
        levels = {}
        for lvl in ab.findall('w:lvl', NS):
            ilvl = lvl.attrib.get(f'{{{NS["w"]}}}ilvl')
            fmt_el = lvl.find('w:numFmt', NS)
            fmt = fmt_el.attrib.get(f'{{{NS["w"]}}}val') if fmt_el is not None else 'decimal'
            txt_el = lvl.find('w:lvlText', NS)
            lvlText = txt_el.attrib.get(f'{{{NS["w"]}}}val') if txt_el is not None else '%1.'
            start_el = lvl.find('w:start', NS)
            start = start_el.attrib.get(f'{{{NS["w"]}}}val') if start_el is not None else '1'
            levels[ilvl] = {'fmt': fmt, 'lvlText': lvlText, 'start': int(start)}
        abstract_nums[ab_id] = levels

    num_to_ab = {}
    for num in root.findall('w:num', NS):
        nid = num.attrib.get(f'{{{NS["w"]}}}numId')
        abid_el = num.find('w:abstractNumId', NS)
        if abid_el is not None:
            num_to_ab[nid] = abid_el.attrib.get(f'{{{NS["w"]}}}val')

    chn = ['零', '一', '二', '三', '四', '五', '六', '七', '八', '九', '十', '十一', '十二', '十三', '十四', '十五']
    counters = {}

    def get_prefix(numPr):
        if numPr is None:
            return None
        nid_elem = numPr.find('w:numId', NS)
        ilvl_elem = numPr.find('w:ilvl', NS)
        if nid_elem is None:
            return None
        nid = nid_elem.attrib.get(f'{{{NS["w"]}}}val')
        if nid == '0' or nid not in num_to_ab:
            return None
        ilvl = ilvl_elem.attrib.get(f'{{{NS["w"]}}}val', '0') if ilvl_elem is not None else '0'

        abid = num_to_ab.get(nid)
        lvl_info = abstract_nums.get(abid, {}).get(ilvl, {})
        fmt = lvl_info.get('fmt', 'decimal')
        lvlText = lvl_info.get('lvlText', '%1.')

        if nid not in counters:
            counters[nid] = {}
        if ilvl not in counters[nid]:
            counters[nid][ilvl] = lvl_info.get('start', 1)
        else:
            counters[nid][ilvl] += 1

        c_val = counters[nid][ilvl]
        if fmt == 'chineseCounting':
            s_val = chn[c_val] if c_val < len(chn) else str(c_val)
        else:
            s_val = str(c_val)

        return lvlText.replace('%1', s_val)

    return get_prefix

def extract_media(docx_path, output_dir):
    """Extract non-wmf media files into output directory and convert EMF to PNG if needed.

    The directory is created lazily, on the first file actually written, so a
    document without images leaves no empty figures_<name>/ behind.
    """
    extracted = []
    emf_files = []
    with zipfile.ZipFile(docx_path) as z:
        for name in z.namelist():
            if name.startswith("word/media/") and not name.endswith("/") and not name.endswith(".wmf"):
                fname = os.path.basename(name)
                if fname:
                    os.makedirs(output_dir, exist_ok=True)
                    target_path = os.path.join(output_dir, fname)
                    with open(target_path, "wb") as f:
                        f.write(z.read(name))
                    extracted.append(fname)
                    if fname.lower().endswith(".emf"):
                        emf_files.append((fname, target_path))

    # Automatically convert EMF files to PNG using headless LibreOffice
    for fname, target_path in emf_files:
        try:
            import subprocess
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "png", target_path, "--outdir", output_dir],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL
            )
            png_name = os.path.splitext(fname)[0] + ".png"
            png_path = os.path.join(output_dir, png_name)
            if os.path.exists(png_path):
                extracted.append(png_name)
                # Auto crop whitespace borders
                try:
                    from PIL import Image, ImageChops
                    im = Image.open(png_path)
                    bg = Image.new(im.mode, im.size, (255, 255, 255))
                    diff = ImageChops.difference(im, bg)
                    bbox = diff.getbbox()
                    if bbox:
                        cropped = im.crop(bbox)
                        cropped.save(png_path)
                except Exception:
                    pass
        except Exception:
            pass

    return extracted

def get_rel_map(docx_path):
    """Map Relationship Id to Target file path."""
    with zipfile.ZipFile(docx_path) as z:
        rels_xml = z.read("word/_rels/document.xml.rels")
    rels_root = ET.fromstring(rels_xml)
    r_ns = {'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}
    return {rel.attrib['Id']: rel.attrib['Target'] for rel in rels_root.findall('.//rel:Relationship', r_ns)}

def extract_ole_number(target_str):
    m = re.search(r'oleObject(\d+)\.bin', target_str)
    return int(m.group(1)) if m else None

def get_node_text(node):
    """Extract plain text from w:t, w:tab, w:br nodes."""
    parts = []
    for elem in node.iter():
        if elem.tag == f"{{{NS['w']}}}t":
            if elem.text:
                parts.append(elem.text)
        elif elem.tag == f"{{{NS['w']}}}tab":
            parts.append("    ")
        elif elem.tag == f"{{{NS['w']}}}br":
            parts.append(" ")
    return "".join(parts)

def parse_run_text_and_math(run_elem, rel_map, formula_map=None, rel_media_dir=None, abs_media_dir=None, image_alts=None):
    """Extract text, MathType OLE formulas, and inline drawings from a run."""
    if formula_map is None:
        formula_map = {}
    if image_alts is None:
        image_alts = {}

    ole = run_elem.find('.//o:OLEObject', NS)
    if ole is not None:
        rid = ole.attrib.get(f"{{{NS['r']}}}id")
        target = rel_map.get(rid, '')
        oid = extract_ole_number(target)
        if oid is not None:
            if oid in formula_map:
                return f" ${formula_map[oid]}$ "
            elif str(oid) in formula_map:
                return f" ${formula_map[str(oid)]}$ "
            else:
                return f" Formula_{oid} "

    # Handle inline drawings inside run
    drawings = run_elem.findall('.//a:blip', NS)
    if drawings:
        inline_imgs = []
        for blip in drawings:
            embed_id = blip.attrib.get(f"{{{NS['r']}}}embed")
            target = rel_map.get(embed_id, '')
            if target and not target.lower().endswith('.wmf'):
                fname = os.path.basename(target)
                if fname.lower().endswith('.emf'):
                    png_cand = os.path.splitext(fname)[0] + '.png'
                    if abs_media_dir and os.path.exists(os.path.join(abs_media_dir, png_cand)):
                        fname = png_cand
                alt = image_alts.get(fname, "")
                media_prefix = f"{rel_media_dir}/" if rel_media_dir else ""
                inline_imgs.append(f" ![{alt}]({media_prefix}{fname}) ")
        if inline_imgs:
            return "".join(inline_imgs)

    t_elems = run_elem.findall('w:t', NS)
    text = "".join([t.text for t in t_elems if t.text])
    if run_elem.find('w:tab', NS) is not None:
        text = "    " + text
    if run_elem.find('w:br', NS) is not None:
        text = text + " "

    if not text:
        return ""

    rPr = run_elem.find('w:rPr', NS)
    is_bold = False
    is_italic = False
    if rPr is not None:
        b_elem = rPr.find('w:b', NS)
        if b_elem is not None and b_elem.attrib.get(f"{{{NS['w']}}}val") != '0':
            is_bold = True
        i_elem = rPr.find('w:i', NS)
        if i_elem is not None and i_elem.attrib.get(f"{{{NS['w']}}}val") != '0':
            is_italic = True

    if is_bold and text.strip():
        return f"**{text.strip()}** "
    if is_italic and text.strip():
        return f"*{text.strip()}* "
    return text

def convert_docx_to_markdown(docx_path, output_md_path, media_dir=None, formula_map=None, image_alts=None):
    """Full-fidelity pipeline for converting Word docx to publication-grade Markdown."""
    if formula_map is None:
        formula_map = {}
    if image_alts is None:
        image_alts = {}

    output_dir = os.path.dirname(os.path.abspath(output_md_path))
    if media_dir is None:
        base = os.path.splitext(os.path.basename(docx_path))[0]
        abs_media_dir = os.path.join(output_dir, f"figures_{base}")
        rel_media_dir = f"figures_{base}"
    else:
        if os.path.isabs(media_dir):
            abs_media_dir = media_dir
            rel_media_dir = os.path.relpath(media_dir, output_dir)
        else:
            abs_media_dir = os.path.join(output_dir, media_dir)
            rel_media_dir = media_dir

    extract_media(docx_path, abs_media_dir)
    rel_map = get_rel_map(docx_path)
    get_prefix = build_numbering_tracker(docx_path)

    doc = docx.Document(docx_path)
    body = doc._element.body

    p_style_map = {}
    for p in doc.paragraphs:
        p_style_map[p._p] = p.style.name if p.style else "Normal"

    tbl_idx_map = {tbl._tbl: i for i, tbl in enumerate(doc.tables)}
    raw_blocks = []

    for element in body:
        if element.tag.endswith('}p'):
            style = p_style_map.get(element, "Normal")
            numPr = element.find('.//w:numPr', NS)
            prefix = get_prefix(numPr)

            # Discover all DrawingML BLIP images in this paragraph with deduplication
            p_images = []
            for blip in element.findall('.//a:blip', NS):
                embed_id = blip.attrib.get(f"{{{NS['r']}}}embed")
                target = rel_map.get(embed_id, '')
                if target and not target.lower().endswith('.wmf'):
                    fname = os.path.basename(target)
                    if fname.lower().endswith('.emf'):
                        png_cand = os.path.splitext(fname)[0] + '.png'
                        if os.path.exists(os.path.join(abs_media_dir, png_cand)):
                            fname = png_cand
                    if fname not in p_images:
                        p_images.append(fname)

            # Check if paragraph contains plain text
            has_plain_text = any(bool(t.text and t.text.strip()) for t in element.findall(f'.//{{{NS["w"]}}}t'))
            if not has_plain_text:
                if p_images:
                    raw_blocks.append(("img_group", p_images))
                continue

            # Parse paragraph runs for paragraphs with text
            runs_parts = []
            for child in element:
                tag = child.tag
                if tag == f"{{{NS['w']}}}r" or tag == f"{{{NS['w']}}}object":
                    runs_parts.append(parse_run_text_and_math(child, rel_map, formula_map, rel_media_dir, abs_media_dir, image_alts))
                elif tag == f"{{{NS['w']}}}hyperlink":
                    hl_id = child.attrib.get(f"{{{NS['r']}}}id")
                    url = rel_map.get(hl_id, "")
                    hl_parts = []
                    for r_child in child.findall(f".//{{{NS['w']}}}r", NS):
                        hl_parts.append(parse_run_text_and_math(r_child, rel_map, formula_map, rel_media_dir, abs_media_dir, image_alts))
                    link_txt = "".join(hl_parts).strip()
                    if url and link_txt:
                        if link_txt == url:
                            runs_parts.append(f"<{url}>")
                        else:
                            runs_parts.append(f"[{link_txt}]({url})")
                    elif link_txt:
                        runs_parts.append(link_txt)
                    elif url:
                        runs_parts.append(f"<{url}>")

            p_text = "".join(runs_parts).strip()
            if not p_text:
                if p_images:
                    raw_blocks.append(("img_group", p_images))
                continue

            clean_txt = clean_inline_text(p_text)
            if not clean_txt:
                if p_images:
                    raw_blocks.append(("img_group", p_images))
                continue

            # Prepend numbering prefix if available
            if prefix and not clean_txt.startswith(prefix):
                sep = "" if prefix.endswith("）") else " "
                clean_txt = f"{prefix}{sep}{clean_txt}"

            # Heading candidate validation guard
            is_heading_candidate = (
                len(clean_txt) <= 45
                and not clean_txt.endswith("。")
                and not clean_txt.endswith("：")
                and not clean_txt.endswith(":")
                and not any(clean_txt.endswith(x) for x in ["如下", "所示", "如下所示", "如下表", "如下图"])
            )

            # Document Title
            is_first_block = (len(raw_blocks) == 0)
            if (style in ("Heading 1", "Title") or is_first_block) and len(clean_txt) <= 40 and not clean_txt.endswith(("。", "：", ":")) and not re.match(r'^[一二三四五六七八九十\d]', clean_txt):
                raw_blocks.append(("h1", f"# {clean_txt.replace('**', '').strip()}\n\n"))
                continue
            # Abstract
            if clean_txt == "摘要" and style == "摘要":
                raw_blocks.append(("h2", "## 摘要\n\n"))
                continue
            # Level 1 section (Chinese Counting: 一、, 二、, etc.)
            if re.match(r'^[一二三四五六七八九十]+、', clean_txt) and is_heading_candidate:
                raw_blocks.append(("h2", f"## {clean_txt.replace('**', '').strip()}\n\n"))
                continue
            # Common section titles (e.g. 实验目标, 实验内容, 实验步骤, etc.)
            clean_title = clean_txt.replace('**', '').rstrip('：: ').strip()
            SECTION_KEYWORDS = (
                "实验目标", "实验目的", "实验内容", "实验步骤", "实验要求",
                "实验原理", "实验环境", "实验仪器", "实验器材", "实验结果",
                "实验总结", "实验报告", "实验思考", "思考题", "操作步骤",
                "注意事项", "背景介绍", "参考文献", "致谢"
            )
            if (clean_title in SECTION_KEYWORDS or any(clean_title == kw for kw in SECTION_KEYWORDS)) and len(clean_title) <= 25 and not clean_txt.endswith("。"):
                raw_blocks.append(("h2", f"## {clean_title}\n\n"))
                continue
            # Subsection 1.1, 5.1, etc.
            if (style == "1.1" or re.match(r'^\d+\.\d+\s+', clean_txt)) and is_heading_candidate:
                raw_blocks.append(("h3", f"### {clean_txt.replace('**', '').strip()}\n\n"))
                continue
            # Subsubsection 5.1.1, etc.
            if (style == "5.1.1" or re.match(r'^\d+\.\d+\.\d+\s+', clean_txt)) and is_heading_candidate:
                raw_blocks.append(("h4", f"#### {clean_txt.replace('**', '').strip()}\n\n"))
                continue
            # Chinese numeral sub-items starting with （一）, （二）, etc.
            if re.match(r'^[（\(][一二三四五六七八九十]+[）\)]', clean_txt) and is_heading_candidate:
                raw_blocks.append(("h3", f"### {clean_txt.replace('**', '').strip()}\n\n"))
                continue
            # Sub-items starting with （1）, （2）, etc. only if style is explicitly a Heading style
            if style in ("Heading 4", "4") and is_heading_candidate:
                raw_blocks.append(("h4", f"#### {clean_txt.replace('**', '').strip()}\n\n"))
                continue
            # Captions
            if style == "图注" or (clean_txt.startswith("图") and any(c.isdigit() for c in clean_txt[:8]) and len(clean_txt) < 120):
                raw_blocks.append(("cap", f"**{clean_txt}**\n\n"))
                continue

            raw_blocks.append(("p", f"{clean_txt}\n\n"))

        elif element.tag.endswith('}tbl'):
            t_idx = tbl_idx_map.get(element)
            if t_idx is None:
                continue
            tbl = doc.tables[t_idx]
            rows = tbl.rows
            if not rows:
                continue

            first_row_cells = rows[0].cells

            # Check if this is an equation table (3 cols, col 3 has equation tag)
            if len(first_row_cells) == 3:
                last_txt = first_row_cells[-1].text.strip()
                if ('（' in last_txt or '(' in last_txt) and any(c.isdigit() for c in last_txt):
                    eq_mds = []
                    for r_idx, row in enumerate(rows):
                        c_last = row.cells[-1].text.strip()
                        m = re.search(r'[\(（]([\d\.\-]+)[\)）]', c_last)
                        tag = m.group(1) if m else c_last.strip('()（）')

                        formula_latex = formula_map.get(tag)
                        if not formula_latex:
                            oles = row.cells[1]._element.findall('.//o:OLEObject', NS)
                            for o in oles:
                                rid = o.attrib.get(f"{{{NS['r']}}}id")
                                oid = extract_ole_number(rel_map.get(rid, ''))
                                if oid is not None:
                                    if oid in formula_map:
                                        formula_latex = formula_map[oid]
                                        break
                                    elif str(oid) in formula_map:
                                        formula_latex = formula_map[str(oid)]
                                        break
                        if not formula_latex:
                            formula_latex = row.cells[1].text.strip()

                        eq_mds.append(f"$$\n{formula_latex} \\tag{{{tag}}}\n$$")
                    raw_blocks.append(("eq_tbl", "\n\n".join(eq_mds) + "\n\n"))
                    continue

            # Check if this is an appendix source code table (1 col)
            if len(first_row_cells) == 1:
                app_title = first_row_cells[0].text.strip()
                if "附录" in app_title and len(rows) > 1:
                    content_cell = rows[1].cells[0]
                    cell_text = content_cell.text.strip()
                    if "import " in cell_text or "def " in cell_text or "class " in cell_text:
                        code_lines = [p.text for p in content_cell.paragraphs]
                        code_body = "\n".join(code_lines).strip()
                        raw_blocks.append(("app_code", f"### {app_title}\n\n```python\n{code_body}\n```\n\n"))
                        continue

            # Regular GFM data table
            table_rows = []
            for r_idx, row in enumerate(rows):
                row_vals = []
                for cell in row.cells:
                    cell_parts = []
                    for p in cell.paragraphs:
                        p_parts = []
                        for child in p._p:
                            tag = child.tag
                            if tag == f"{{{NS['w']}}}r" or tag == f"{{{NS['w']}}}object":
                                p_parts.append(parse_run_text_and_math(child, rel_map, formula_map))
                        cell_parts.append("".join(p_parts).strip())
                    val = " ".join([part for part in cell_parts if part]).replace('|', '\\|').replace('\n', ' ')
                    val = clean_inline_text(val)
                    row_vals.append(val)

                # Skip completely empty rows or orphan symbol rows
                if all(not v for v in row_vals):
                    continue
                if len(row_vals) == 2 and not row_vals[1]:
                    continue

                table_rows.append("| " + " | ".join(row_vals) + " |")
                if len(table_rows) == 1:
                    table_rows.append("| " + " | ".join([":---" if ci == 0 else ":---:" for ci in range(len(row_vals))]) + " |")

            if table_rows:
                raw_blocks.append(("data_tbl", "\n".join(table_rows) + "\n\n"))

    # Post-process blocks: Image & Caption Binding & Side-by-side unbundling
    final_output = []
    i = 0
    while i < len(raw_blocks):
        btype, bdata = raw_blocks[i]

        if btype == "img_group":
            img_list = bdata
            # Look ahead for caption
            caption = None
            cap_idx = None
            for j in range(1, 4):
                if i + j < len(raw_blocks):
                    nxt_type, nxt_content = raw_blocks[i + j]
                    if not isinstance(nxt_content, str):
                        continue
                    nxt_clean = nxt_content.strip('*_ \n')
                    if (nxt_clean.startswith("图") or nxt_clean.startswith("Figure")) and len(nxt_clean) < 120:
                        caption = nxt_clean
                        cap_idx = i + j
                        break

            if cap_idx is not None:
                raw_blocks[cap_idx] = ("skip", "")

            # Check if images have distinct figure numbers
            fig_nums = []
            for img_name in img_list:
                alt = image_alts.get(img_name, "")
                m_fig = re.match(r'^(图[\d\.\-]+)', alt)
                fig_nums.append(m_fig.group(1) if m_fig else alt)

            # Check if each image has a distinct alt in image_alts
            distinct_alts = [image_alts.get(img_name, "") for img_name in img_list]
            has_distinct_alts = len(img_list) > 1 and all(distinct_alts) and len(set(distinct_alts)) == len(img_list)

            if (len(img_list) > 1 and len(set(fig_nums)) == len(img_list) and all(bool(f) for f in fig_nums)) or has_distinct_alts:
                # Distinct figures placed side-by-side in Word: render each image with its own caption
                blocks = []
                for img_name in img_list:
                    alt = image_alts.get(img_name, "")
                    if alt:
                        blocks.append(f"![{alt}]({rel_media_dir}/{img_name})\n\n**{alt}**")
                    else:
                        blocks.append(f"![]({rel_media_dir}/{img_name})")
                final_output.append("\n\n".join(blocks) + "\n\n")
            else:
                # Subfigures sharing the same figure number or single figure
                first_img = img_list[0] if img_list else ""
                default_cap = caption if caption else (image_alts.get(first_img, "") if first_img else "")
                img_mds = []
                for img_name in img_list:
                    alt = image_alts.get(img_name, default_cap or "")
                    img_mds.append(f"![{alt}]({rel_media_dir}/{img_name})")
                if caption:
                    final_output.append("\n\n".join(img_mds) + f"\n\n**{caption}**\n\n")
                elif default_cap:
                    final_output.append("\n\n".join(img_mds) + f"\n\n**{default_cap}**\n\n")
                else:
                    final_output.append("\n\n".join(img_mds) + "\n\n")
            i += 1
            continue

        if btype == "skip":
            i += 1
            continue

        final_output.append(bdata)
        i += 1

    result_text = "".join(final_output)
    result_text = re.sub(r'\n{3,}', '\n\n', result_text)
    result_text = re.sub(r'(\[\d+\])\s+', r'\1 ', result_text)

    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write(result_text)

    print(f"[+] Successfully converted '{docx_path}' -> '{output_md_path}' ({len(result_text):,} chars)")

def main():
    parser = argparse.ArgumentParser(description="High-Fidelity DOCX to Markdown Converter with Universal Formula Mapping.")
    parser.add_argument("input_docx", help="Path to input .docx file")
    parser.add_argument("output_md", nargs="?", default=None, help="Path to output .md file (default: same basename as docx)")
    parser.add_argument("--media-dir", default=None, help="Directory to save extracted figures (default: figures_<docname>)")
    parser.add_argument("--formula-map", default=None, help="JSON file with equation mappings (tag-based or OLE-based)")
    parser.add_argument("--image-alts", default=None, help="JSON file with mapping of image filenames to captions/alt texts")
    args = parser.parse_args()

    input_file = args.input_docx
    if not os.path.exists(input_file):
        print(f"Error: File not found: {input_file}", file=sys.stderr)
        sys.exit(1)

    # Automatically convert legacy .doc to .docx using headless LibreOffice if needed
    if input_file.lower().endswith('.doc') and not input_file.lower().endswith('.docx'):
        docx_file = os.path.splitext(input_file)[0] + ".docx"
        out_dir = os.path.dirname(os.path.abspath(input_file))
        print(f"[*] Detected legacy .doc file. Converting to .docx via LibreOffice...")
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "docx", input_file, "--outdir", out_dir],
            check=True
        )
        input_docx = docx_file
    else:
        input_docx = input_file

    output_md = args.output_md or os.path.splitext(input_file)[0] + ".md"

    formula_map = {}
    if args.formula_map:
        if os.path.exists(args.formula_map):
            with open(args.formula_map, "r", encoding="utf-8") as f:
                loaded = json.load(f)
                # Allow both integer and string keys
                for k, v in loaded.items():
                    formula_map[k] = v
                    if k.isdigit():
                        formula_map[int(k)] = v
        else:
            print(f"Warning: Formula map file '{args.formula_map}' not found.", file=sys.stderr)

    image_alts = {}
    if args.image_alts:
        if os.path.exists(args.image_alts):
            with open(args.image_alts, "r", encoding="utf-8") as f:
                image_alts = json.load(f)
        else:
            print(f"Warning: Image alts file '{args.image_alts}' not found.", file=sys.stderr)

    convert_docx_to_markdown(
        input_docx,
        output_md,
        media_dir=args.media_dir,
        formula_map=formula_map,
        image_alts=image_alts
    )

if __name__ == '__main__':
    main()
